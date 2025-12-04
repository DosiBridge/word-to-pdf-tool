"""
Comprehensive test script for all backend API endpoints.
Tests: PDF to Word, Word to PDF, PDF to Text, Unlock PDF, Merge PDFs, Split PDF
"""
import requests
import os
from reportlab.pdfgen import canvas
from docx import Document
import time

BASE_URL = "http://localhost:8000/api"
RESULTS = []

def log_result(name, success, message=""):
    status = "✅ PASS" if success else "❌ FAIL"
    RESULTS.append((name, success, message))
    print(f"{status}: {name}" + (f" - {message}" if message else ""))

def create_test_files():
    print("\n=== Creating Test Files ===")
    
    # Create multi-page PDF
    c = canvas.Canvas("test_multipage.pdf")
    for i in range(1, 4):
        c.drawString(100, 750, f"This is page {i} of the test PDF")
        c.drawString(100, 700, f"Content for page {i}")
        c.showPage()
    c.save()
    print("Created: test_multipage.pdf (3 pages)")
    
    # Create simple PDF for conversions
    c = canvas.Canvas("test_simple.pdf")
    c.drawString(100, 750, "Hello World - PDF to Word Test")
    c.drawString(100, 700, "This text should appear in the Word document")
    c.save()
    print("Created: test_simple.pdf")

    # Create Word document
    doc = Document()
    doc.add_heading("Test Word Document", 0)
    doc.add_paragraph("This is a test paragraph for Word to PDF conversion.")
    doc.add_paragraph("Second paragraph with more content.")
    doc.save("test_document.docx")
    print("Created: test_document.docx")
    
    print("Test files created successfully!\n")

def test_health():
    print("\n=== Testing Server Health ===")
    try:
        response = requests.get("http://localhost:8000/")
        if response.status_code == 200:
            log_result("Server Health", True, "Server is running")
            return True
        else:
            log_result("Server Health", False, f"Status: {response.status_code}")
            return False
    except Exception as e:
        log_result("Server Health", False, str(e))
        return False

def test_pdf_to_word():
    print("\n=== Testing PDF to Word ===")
    try:
        with open("test_simple.pdf", "rb") as f:
            files = {"file": ("test_simple.pdf", f, "application/pdf")}
            response = requests.post(f"{BASE_URL}/pdf-to-word", files=files)
        
        if response.status_code == 200:
            with open("output_pdf_to_word.docx", "wb") as out:
                out.write(response.content)
            size = os.path.getsize("output_pdf_to_word.docx")
            log_result("PDF to Word", True, f"Output size: {size} bytes")
            return True
        else:
            log_result("PDF to Word", False, f"Status: {response.status_code}, Response: {response.text[:200]}")
            return False
    except Exception as e:
        log_result("PDF to Word", False, str(e))
        return False

def test_word_to_pdf():
    print("\n=== Testing Word to PDF ===")
    try:
        with open("test_document.docx", "rb") as f:
            files = {"file": ("test_document.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
            response = requests.post(f"{BASE_URL}/word-to-pdf", files=files)
        
        if response.status_code == 200:
            with open("output_word_to_pdf.pdf", "wb") as out:
                out.write(response.content)
            size = os.path.getsize("output_word_to_pdf.pdf")
            log_result("Word to PDF", True, f"Output size: {size} bytes")
            return True
        else:
            log_result("Word to PDF", False, f"Status: {response.status_code}, Response: {response.text[:200]}")
            return False
    except Exception as e:
        log_result("Word to PDF", False, str(e))
        return False

def test_pdf_to_text():
    print("\n=== Testing PDF to Text ===")
    try:
        with open("test_simple.pdf", "rb") as f:
            files = {"file": ("test_simple.pdf", f, "application/pdf")}
            response = requests.post(f"{BASE_URL}/pdf-to-txt", files=files)
        
        if response.status_code == 200:
            text = response.text
            if "Hello World" in text or len(text) > 0:
                log_result("PDF to Text", True, f"Extracted {len(text)} characters")
                return True
            else:
                log_result("PDF to Text", False, "No text extracted")
                return False
        else:
            log_result("PDF to Text", False, f"Status: {response.status_code}, Response: {response.text[:200]}")
            return False
    except Exception as e:
        log_result("PDF to Text", False, str(e))
        return False

def test_unlock_pdf():
    print("\n=== Testing Unlock PDF ===")
    try:
        with open("test_simple.pdf", "rb") as f:
            files = {"file": ("test_simple.pdf", f, "application/pdf")}
            response = requests.post(f"{BASE_URL}/pdf-unlock", files=files)
        
        if response.status_code == 200:
            with open("output_unlocked.pdf", "wb") as out:
                out.write(response.content)
            size = os.path.getsize("output_unlocked.pdf")
            log_result("Unlock PDF", True, f"Output size: {size} bytes")
            return True
        else:
            log_result("Unlock PDF", False, f"Status: {response.status_code}, Response: {response.text[:200]}")
            return False
    except Exception as e:
        log_result("Unlock PDF", False, str(e))
        return False

def test_merge_pdfs():
    print("\n=== Testing Merge PDFs ===")
    try:
        with open("test_simple.pdf", "rb") as f1, open("test_multipage.pdf", "rb") as f2:
            files = [
                ("files", ("file1.pdf", f1, "application/pdf")),
                ("files", ("file2.pdf", f2, "application/pdf"))
            ]
            response = requests.post(f"{BASE_URL}/merge-pdfs", files=files)
        
        if response.status_code == 200:
            with open("output_merged.pdf", "wb") as out:
                out.write(response.content)
            size = os.path.getsize("output_merged.pdf")
            log_result("Merge PDFs", True, f"Output size: {size} bytes")
            return True
        else:
            log_result("Merge PDFs", False, f"Status: {response.status_code}, Response: {response.text[:200]}")
            return False
    except Exception as e:
        log_result("Merge PDFs", False, str(e))
        return False

def test_split_pdf():
    print("\n=== Testing Split PDF ===")
    try:
        with open("test_multipage.pdf", "rb") as f:
            files = {"file": ("test_multipage.pdf", f, "application/pdf")}
            data = {"pages": "1,2"}
            response = requests.post(f"{BASE_URL}/split-pdf", files=files, data=data)
        
        if response.status_code == 200:
            with open("output_split.pdf", "wb") as out:
                out.write(response.content)
            size = os.path.getsize("output_split.pdf")
            log_result("Split PDF", True, f"Output size: {size} bytes")
            return True
        else:
            log_result("Split PDF", False, f"Status: {response.status_code}, Response: {response.text[:200]}")
            return False
    except Exception as e:
        log_result("Split PDF", False, str(e))
        return False

def cleanup():
    print("\n=== Cleaning Up ===")
    test_files = [
        "test_multipage.pdf", "test_simple.pdf", "test_document.docx",
        "output_pdf_to_word.docx", "output_word_to_pdf.pdf", "output_unlocked.pdf",
        "output_merged.pdf", "output_split.pdf"
    ]
    for f in test_files:
        if os.path.exists(f):
            os.remove(f)
            print(f"Removed: {f}")

def print_summary():
    print("\n" + "="*50)
    print("           TEST RESULTS SUMMARY")
    print("="*50)
    passed = sum(1 for _, success, _ in RESULTS if success)
    failed = sum(1 for _, success, _ in RESULTS if not success)
    
    for name, success, message in RESULTS:
        status = "✅" if success else "❌"
        print(f"  {status} {name}")
    
    print("="*50)
    print(f"  Total: {len(RESULTS)} | Passed: {passed} | Failed: {failed}")
    print("="*50)
    
    if failed == 0:
        print("\n🎉 ALL TESTS PASSED!")
    else:
        print(f"\n⚠️  {failed} TEST(S) FAILED - Review errors above")

if __name__ == "__main__":
    print("="*50)
    print("    COMPREHENSIVE API ENDPOINT TESTING")
    print("="*50)
    
    try:
        create_test_files()
        
        if not test_health():
            print("\n❌ Server is not running! Please start the backend server.")
            exit(1)
        
        test_pdf_to_word()
        test_word_to_pdf()
        test_pdf_to_text()
        test_unlock_pdf()
        test_merge_pdfs()
        test_split_pdf()
        
    except Exception as e:
        print(f"\n❌ Test suite error: {e}")
        import traceback
        traceback.print_exc()
    finally:
        cleanup()
        print_summary()
